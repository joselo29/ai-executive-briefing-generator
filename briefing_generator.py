"""AI Executive Briefing Generator: HR CSV -> Gemini (schema + per-section) -> .docx."""

from __future__ import annotations

import json
import os
import re
import sys
import time
from datetime import date
from pathlib import Path

import pandas as pd
from google import genai
from google.genai import errors as genai_errors
from google.genai import types
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parent
DATA_PATH = ROOT / "data" / "sample_input.csv"
PROMPTS_PATH = ROOT / "prompts.md"
OUTPUT_DIR = ROOT / "output"
OUTPUT_PATH = OUTPUT_DIR / "sample_briefing.docx"
MODEL_NAME = "gemini-2.5-flash"


def load_prompts(path: Path = PROMPTS_PATH) -> dict[str, str]:
    """Parse prompts.md into a {PROMPT_N: body} dict, stripping inline comments."""
    text = path.read_text(encoding="utf-8")
    blocks = re.split(r"^## PROMPT_", text, flags=re.MULTILINE)[1:]
    prompts: dict[str, str] = {}
    for block in blocks:
        lines = block.splitlines()
        key = "PROMPT_" + lines[0].strip().split()[0]
        body_lines = [
            line for line in lines[1:]
            if not line.strip().startswith("#") and line.strip() != "---"
        ]
        prompts[key] = "\n".join(body_lines).strip()
    return prompts


def data_loader(path: Path = DATA_PATH) -> pd.DataFrame:
    """Load the input CSV and normalise column names."""
    df = pd.read_csv(path)
    df.columns = df.columns.str.strip()
    return df


def _configure_gemini() -> genai.Client:
    load_dotenv(ROOT / ".env")
    api_key = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
    if not api_key:
        raise RuntimeError("GEMINI_API_KEY missing from .env")
    return genai.Client(api_key=api_key)


_RETRY_WAITS = (5, 10, 20)
_RETRY_CODES = {429, 503}


def _call_gemini(client: genai.Client, prompt: str, **kwargs) -> str:
    attempt = 0
    while True:
        try:
            response = client.models.generate_content(
                model=MODEL_NAME, contents=prompt, **kwargs
            )
            return (response.text or "").strip()
        except (genai_errors.ServerError, genai_errors.ClientError) as exc:
            code = getattr(exc, "code", None)
            if code not in _RETRY_CODES or attempt >= len(_RETRY_WAITS):
                raise
            wait = _RETRY_WAITS[attempt]
            attempt += 1
            print(
                f"[retry] Gemini unavailable (code {code}) — waiting {wait}s and "
                f"retrying (attempt {attempt}/3)...",
                file=sys.stderr,
            )
            time.sleep(wait)


_JSON_FENCE = re.compile(r"```(?:json)?\s*|\s*```", re.MULTILINE)


def _extract_json(raw: str) -> dict:
    cleaned = _JSON_FENCE.sub("", raw).strip()
    match = re.search(r"\{.*\}", cleaned, re.DOTALL)
    if not match:
        raise json.JSONDecodeError("No JSON object found", cleaned, 0)
    return json.loads(match.group(0))


def schema_discovery(
    df: pd.DataFrame,
    client: genai.Client,
    prompts: dict[str, str],
) -> dict:
    """Call 1: ask Gemini to return JSON describing dimensions + recommended sections."""
    schema_context = (
        "Columns and dtypes:\n"
        + "\n".join(f"- {c}: {t}" for c, t in df.dtypes.items())
        + "\n\nSample rows (3):\n"
        + df.head(3).to_csv(index=False)
    )
    prompt = prompts["PROMPT_1"] + "\n\n" + schema_context

    last_error: Exception | None = None
    for attempt in range(2):
        try:
            raw = _call_gemini(
                client,
                prompt,
                config=types.GenerateContentConfig(response_mime_type="application/json"),
            )
            data = _extract_json(raw)
            if "available_dimensions" not in data or "recommended_sections" not in data:
                raise ValueError("Missing required keys in schema JSON")
            return data
        except (json.JSONDecodeError, ValueError) as exc:
            last_error = exc
            print(
                f"[schema_discovery] attempt {attempt + 1} failed: {exc}",
                file=sys.stderr,
            )
    raise RuntimeError(f"Schema discovery failed after retry: {last_error}")


def _executive_summary_data(df: pd.DataFrame) -> str:
    total = len(df)
    attr_yes = int((df["Attrition"] == "Yes").sum())
    attr_rate = attr_yes / total * 100
    return (
        f"Total employees: {total}\n"
        f"Attrition (left): {attr_yes} ({attr_rate:.1f}%)\n"
        f"Average JobSatisfaction (1-4 scale): {df['JobSatisfaction'].mean():.2f}\n"
        f"Average MonthlyIncome: ${df['MonthlyIncome'].mean():,.0f}\n"
        f"Average tenure (YearsAtCompany): {df['YearsAtCompany'].mean():.1f} years\n"
        f"Employees working OverTime: {(df['OverTime'] == 'Yes').mean() * 100:.1f}%\n"
        f"Headcount by department: {df['Department'].value_counts().to_dict()}\n"
    )


def _attrition_data(df: pd.DataFrame) -> str:
    by_dept = df.groupby("Department")["Attrition"].apply(
        lambda s: f"{(s == 'Yes').sum()} of {len(s)} ({(s == 'Yes').mean() * 100:.1f}%)"
    ).to_dict()
    by_overtime = df.groupby("OverTime")["Attrition"].apply(
        lambda s: f"{(s == 'Yes').sum()} of {len(s)} ({(s == 'Yes').mean() * 100:.1f}%)"
    ).to_dict()
    by_role = (
        df.groupby("JobRole")["Attrition"]
        .apply(lambda s: round((s == "Yes").mean() * 100, 1))
        .sort_values(ascending=False)
        .head(5)
        .to_dict()
    )
    tenure = df.copy()
    tenure["TenureBucket"] = pd.cut(
        tenure["YearsAtCompany"],
        bins=[-1, 1, 3, 5, 10, 100],
        labels=["0-1", "2-3", "4-5", "6-10", "10+"],
    )
    by_tenure = tenure.groupby("TenureBucket", observed=True)["Attrition"].apply(
        lambda s: f"{(s == 'Yes').mean() * 100:.1f}%"
    ).to_dict()
    return (
        f"Attrition rate by Department: {by_dept}\n"
        f"Attrition rate by OverTime status: {by_overtime}\n"
        f"Top 5 JobRoles by attrition rate (%): {by_role}\n"
        f"Attrition rate by tenure bucket: {by_tenure}\n"
    )


def _satisfaction_data(df: pd.DataFrame) -> str:
    avg_by_dept = df.groupby("Department")["JobSatisfaction"].mean().round(2).to_dict()
    env_by_dept = df.groupby("Department")["EnvironmentSatisfaction"].mean().round(2).to_dict()
    wlb_by_dept = df.groupby("Department")["WorkLifeBalance"].mean().round(2).to_dict()
    low_sat_share = (df["JobSatisfaction"] <= 2).mean() * 100
    low_sat_attr = df[df["JobSatisfaction"] <= 2]["Attrition"].eq("Yes").mean() * 100
    high_sat_attr = df[df["JobSatisfaction"] >= 3]["Attrition"].eq("Yes").mean() * 100
    return (
        f"Average JobSatisfaction by Department (1-4): {avg_by_dept}\n"
        f"Average EnvironmentSatisfaction by Department (1-4): {env_by_dept}\n"
        f"Average WorkLifeBalance by Department (1-4): {wlb_by_dept}\n"
        f"Share of employees with low JobSatisfaction (<=2): {low_sat_share:.1f}%\n"
        f"Attrition rate among low-satisfaction employees: {low_sat_attr:.1f}%\n"
        f"Attrition rate among satisfied employees (>=3): {high_sat_attr:.1f}%\n"
    )


def _department_data(df: pd.DataFrame) -> str:
    perf = df.groupby("Department")["PerformanceRating"].mean().round(2).to_dict()
    income = df.groupby("Department")["MonthlyIncome"].mean().round(0).to_dict()
    attr = df.groupby("Department")["Attrition"].apply(
        lambda s: round((s == "Yes").mean() * 100, 1)
    ).to_dict()
    sat = df.groupby("Department")["JobSatisfaction"].mean().round(2).to_dict()
    return (
        f"Headcount by Department: {df['Department'].value_counts().to_dict()}\n"
        f"Average PerformanceRating by Department (1-4): {perf}\n"
        f"Average MonthlyIncome by Department: {income}\n"
        f"Average JobSatisfaction by Department (1-4): {sat}\n"
        f"Attrition rate by Department (%): {attr}\n"
    )


def _risk_data(df: pd.DataFrame) -> str:
    total = len(df)
    overall_attr = (df["Attrition"] == "Yes").mean() * 100
    overtime_attr = df[df["OverTime"] == "Yes"]["Attrition"].eq("Yes").mean() * 100
    low_sat_attr = df[df["JobSatisfaction"] <= 2]["Attrition"].eq("Yes").mean() * 100
    early_tenure_attr = df[df["YearsAtCompany"] <= 2]["Attrition"].eq("Yes").mean() * 100
    low_income_attr = (
        df[df["MonthlyIncome"] < df["MonthlyIncome"].median()]["Attrition"].eq("Yes").mean() * 100
    )
    role_attr = (
        df.groupby("JobRole")["Attrition"]
        .apply(lambda s: round((s == "Yes").mean() * 100, 1))
        .sort_values(ascending=False)
        .head(3)
        .to_dict()
    )
    return (
        f"Total workforce: {total}\n"
        f"Overall attrition rate: {overall_attr:.1f}%\n"
        f"Attrition among employees working OverTime: {overtime_attr:.1f}%\n"
        f"Attrition among employees with low JobSatisfaction (<=2): {low_sat_attr:.1f}%\n"
        f"Attrition among employees with <=2 years tenure: {early_tenure_attr:.1f}%\n"
        f"Attrition among below-median earners: {low_income_attr:.1f}%\n"
        f"Top 3 highest-attrition job roles (%): {role_attr}\n"
    )


def _compensation_data(df: pd.DataFrame) -> str:
    income_by_dept = df.groupby("Department")["MonthlyIncome"].mean().round(0).to_dict()
    role_income = (
        df.groupby("JobRole")["MonthlyIncome"].mean().round(0).sort_values(ascending=False)
    )
    top_roles = role_income.head(5).to_dict()
    bottom_roles = role_income.tail(5).to_dict()
    income_by_perf = (
        df.groupby("PerformanceRating")["MonthlyIncome"].mean().round(0).to_dict()
    )
    income_by_overtime = (
        df.groupby("OverTime")["MonthlyIncome"].mean().round(0).to_dict()
    )
    return (
        f"Average MonthlyIncome by Department: {income_by_dept}\n"
        f"Top 5 JobRoles by average MonthlyIncome: {top_roles}\n"
        f"Bottom 5 JobRoles by average MonthlyIncome: {bottom_roles}\n"
        f"Average MonthlyIncome by PerformanceRating (1-4): {income_by_perf}\n"
        f"Average MonthlyIncome by OverTime status: {income_by_overtime}\n"
    )


SECTION_BUILDERS = {
    "Executive Summary": ("PROMPT_2", "executive_summary_data", _executive_summary_data),
    "Attrition Analysis": ("PROMPT_3", "attrition_data", _attrition_data),
    "Satisfaction & Engagement": ("PROMPT_4", "satisfaction_data", _satisfaction_data),
    "Department Performance": ("PROMPT_5", "department_data", _department_data),
    "Risks & Recommendations": ("PROMPT_6", "risk_data", _risk_data),
    "Compensation & Reward Strategy": ("PROMPT_9", "compensation_data", _compensation_data),
}


def _match_section(name: str) -> str | None:
    n = name.lower()
    if "exec" in n or "summary" in n:
        return "Executive Summary"
    if "attrition" in n or "turnover" in n:
        return "Attrition Analysis"
    if "satisf" in n or "engage" in n or "morale" in n:
        return "Satisfaction & Engagement"
    if "compensation" in n or "reward" in n or "pay" in n or "salary" in n:
        return "Compensation & Reward Strategy"
    if "department" in n or "performance" in n:
        return "Department Performance"
    if "risk" in n or "recommend" in n:
        return "Risks & Recommendations"
    return None


def generate_section(
    section_name: str,
    df: pd.DataFrame,
    client: genai.Client,
    prompts: dict[str, str],
) -> str:
    """Call 2: build the data slice for a section and ask Gemini for the narrative."""
    matched = _match_section(section_name) or "Executive Summary"
    prompt_key, placeholder, builder = SECTION_BUILDERS[matched]
    data_text = builder(df)
    prompt = prompts[prompt_key].replace("{" + placeholder + "}", data_text)
    return _call_gemini(client, prompt)


def generate_executive_commentary(client: genai.Client, briefing_text: str) -> str:
    """Final pass: ask Gemini for a Chief HR Officer-level commentary on the assembled briefing."""
    prompts = load_prompts()
    prompt = prompts["PROMPT_7"].replace("{briefing_text}", briefing_text)
    return _call_gemini(client, prompt)


def qa_mode(client: genai.Client, briefing_text: str) -> None:
    """Interactive RAG Q&A loop — answers grounded in the briefing context only."""
    prompts = load_prompts()
    template = prompts["PROMPT_8"]
    print("[Q&A] Briefing ready. Ask questions about this report (type 'exit' to quit):")
    while True:
        try:
            question = input("> ").strip()
        except (EOFError, KeyboardInterrupt):
            print()
            return
        if not question:
            continue
        if question.lower() == "exit":
            return
        prompt = template.replace("{briefing_text}", briefing_text).replace(
            "{user_question}", question
        )
        answer = _call_gemini(client, prompt)
        print(answer)
        print()


def build_document(
    sections: list[tuple[str, str]],
    output_path: Path = OUTPUT_PATH,
) -> Path:
    """Assemble the final .docx with title, date, and one heading per section."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(11)

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run("Workforce Executive Briefing")
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run(f"Prepared {date.today().strftime('%B %d, %Y')}")
    sub_run.italic = True
    sub_run.font.size = Pt(11)
    sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph()

    for name, narrative in sections:
        heading = doc.add_heading(name, level=1)
        for run in heading.runs:
            run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        for paragraph in [p.strip() for p in narrative.split("\n\n") if p.strip()]:
            doc.add_paragraph(paragraph)

    doc.save(output_path)
    return output_path


def main() -> None:
    prompts = load_prompts()
    df = data_loader()
    client = _configure_gemini()

    schema = schema_discovery(df, client, prompts)
    print(f"[schema] dimensions: {schema['available_dimensions']}")
    print(f"[schema] sections:   {schema['recommended_sections']}")

    sections: list[tuple[str, str]] = []
    for section_name in schema["recommended_sections"]:
        print(f"[generating] {section_name}")
        narrative = generate_section(section_name, df, client, prompts)
        sections.append((section_name, narrative))

    print("[commentary] Generating Strategic Executive Commentary...")
    briefing_text = "\n\n".join(f"## {name}\n\n{narrative}" for name, narrative in sections)
    commentary = generate_executive_commentary(client, briefing_text)
    sections.append(("Strategic Executive Commentary", commentary))

    out = build_document(sections)
    print(f"[done] wrote {out}")

    if "--qa" in sys.argv[1:]:
        full_briefing = "\n\n".join(f"## {name}\n\n{narrative}" for name, narrative in sections)
        qa_mode(client, full_briefing)


if __name__ == "__main__":
    main()
